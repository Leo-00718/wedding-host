# -*- coding: utf-8 -*-
"""
用原始 Word 模板填内容 —— 生成流程表（1:1 保留 logo、底色、表格样式）。

用法：
    python fill_flow.py <模板.docx> <输出.docx> <替换配置.json>

替换配置是一个 JSON，例如：
{
  "上一场新人名": "新 郎 & 新 娘",
  "12:08": "11:58",
  "上一场酒店 - 旧厅名": "新酒店 - 新厅名"
}

原理：打开模板 docx → 在所有段落和表格单元格里做文本替换 → 落款自动贴页底 → 另存。
设计元素（logo、图片、表格底色、边框）原封不动。
"""
import sys, json, shutil, os, tempfile
from docx import Document
from docx.oxml.ns import qn

WD_ACTIVE_END_PAGE = 3   # Word: Range.Information(wdActiveEndPageNumber)


def replace_in_paragraph(para, repl):
    for run in para.runs:
        for old, new in repl.items():
            if old in run.text:
                run.text = run.text.replace(old, new)


def _place_signature_at_bottom(path):
    """把「logo + 联系电话 + 署名」推到最后一页的底部。

    先删多余空行（已经跨页时），再加空行，直到落款不再和正文同页，再撤回最后一行。
    得到的是「再加一行就翻页」的极限位置，而且每场内容长短不同也能自动适配。

    硬约束：落款整块**不能跨页**。内容太满、两个要求冲突时，优先保证不跨页。

    依赖本机安装 Microsoft Word（pywin32）。没装就跳过，落款会紧跟表格。
    """
    try:
        import win32com.client as wc
    except ImportError:
        print("[警告] 没装 pywin32，跳过落款置底（落款将紧跟表格）")
        return

    try:
        app = wc.DispatchEx("Word.Application")
    except Exception as e:
        print("[警告] 打不开 Word，跳过落款置底（原因：%s）" % e)
        return

    app.Visible = False
    app.DisplayAlerts = 0
    try:
        doc = app.Documents.Open(path)
        try:
            total = doc.Paragraphs.Count
            logo_i = sign_i = None
            for i in range(1, total + 1):
                r = doc.Paragraphs(i).Range
                if "品山主持工作室" in r.Text:
                    sign_i = i
                if logo_i is None and i > total / 2 and r.InlineShapes.Count:
                    logo_i = i
            if not logo_i or not sign_i:
                print("[警告] 没找到落款块，跳过落款置底")
                return

            content_i = None
            for i in range(logo_i - 1, 0, -1):
                if doc.Paragraphs(i).Range.Text.strip():
                    content_i = i
                    break
            if not content_i:
                print("[警告] 没找到正文终点，跳过落款置底")
                return
            end_page = doc.Paragraphs(content_i).Range.Information(WD_ACTIVE_END_PAGE)

            def same_page():
                return (doc.Paragraphs(logo_i).Range.Information(WD_ACTIVE_END_PAGE) == end_page
                        and doc.Paragraphs(sign_i).Range.Information(WD_ACTIVE_END_PAGE) == end_page)

            def blank_at(i):
                r = doc.Paragraphs(i).Range
                return (not r.Text.replace("\r", "").replace("\x07", "").strip()
                        and r.InlineShapes.Count == 0)

            # 1) 已经跨页 —— 从落款往上找最近的可删空行，删到能放进一页为止
            #    （落款块内部的空行在 logo 之后，不会被删到；最多删 5 个，避免吃掉页眉间距）
            removed = 0
            while not same_page() and removed < 5:
                target = None
                for i in range(logo_i - 1, 0, -1):
                    if blank_at(i):
                        target = i
                        break
                if target is None:
                    break
                doc.Paragraphs(target).Range.Delete()
                logo_i -= 1
                sign_i -= 1
                removed += 1

            # 2) 还能往下推 —— 一行一行加，加到刚好不翻页
            added = 0
            while added < 60 and same_page():
                doc.Paragraphs(logo_i).Range.InsertParagraphBefore()
                logo_i += 1
                sign_i += 1
                if same_page():
                    added += 1
                else:
                    doc.Paragraphs(logo_i - 1).Range.Delete()
                    logo_i -= 1
                    sign_i -= 1
                    break

            doc.Save()
            if same_page():
                print("落款已贴页底：加空行 %d、删空行 %d，落在第 %d 页" % (added, removed, end_page))
            else:
                print("[注意] 内容太满，落款无法贴页底，只能紧跟正文（加 %d、删 %d）" % (added, removed))
        finally:
            doc.Close(0)
    finally:
        app.Quit()


def fill(template_path, output_path, repl):
    # 先复制到临时文件再打开，避免"边读边写同一路径"在 Windows 上锁冲突
    tmpdir = tempfile.mkdtemp()
    tmp = os.path.join(tmpdir, "template.docx")
    shutil.copy2(template_path, tmp)
    doc = Document(tmp)

    # 替换段落
    for para in doc.paragraphs:
        replace_in_paragraph(para, repl)

    # 替换表格
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    replace_in_paragraph(para, repl)

    # 替换页眉页脚（如有）
    for section in doc.sections:
        for para in section.header.paragraphs + section.footer.paragraphs:
            replace_in_paragraph(para, repl)

    # 落款整块 keep_together / keep_with_next —— 硬约束：整块不跨页
    paras = doc.paragraphs
    for i, para in enumerate(paras):
        t = para.text.strip()
        if "联系电话" in t:
            para.paragraph_format.keep_together = True
            para.paragraph_format.keep_with_next = True
            if i + 1 < len(paras):
                paras[i + 1].paragraph_format.keep_with_next = True
            # logo 段落也要往下绑，否则会被单独留在上一页
            if i - 1 >= 0 and paras[i - 1]._p.findall(".//" + qn("w:drawing")):
                paras[i - 1].paragraph_format.keep_with_next = True
        elif "品山主持工作室" in t:
            para.paragraph_format.keep_together = True

    doc.save(output_path)

    # 再让 Word 量真实分页，把落款推到「再加一行就翻页」的极限位置
    _place_signature_at_bottom(output_path)


def main():
    if len(sys.argv) < 4:
        print(__doc__)
        return 1
    template, output, cfg = sys.argv[1], sys.argv[2], sys.argv[3]
    repl = json.load(open(cfg, encoding="utf-8"))
    fill(template, output, repl)
    print("完成:", output)
    print("替换了 %d 组文本" % len(repl))
    return 0


if __name__ == "__main__":
    sys.exit(main())